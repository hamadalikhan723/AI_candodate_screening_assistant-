
import os
import re
import time

import pymupdf
import numpy as np

from docx import Document
from sentence_transformers import SentenceTransformer


# ============================================================
# CONFIGURATION
# ============================================================

EMBEDDING_MODEL = "BAAI/bge-small-en-v1.5"
DEVICE = "cpu"

import torch


# Hybrid scoring weights
SEMANTIC_WEIGHT = 0.50
KEYWORD_WEIGHT = 0.30
EVIDENCE_WEIGHT = 0.20

STRONG_MATCH = 0.75
PARTIAL_MATCH = 0.50


# ============================================================
# LOAD EMBEDDING MODEL
# ============================================================

print("=" * 60)
print("Loading BGE embedding model...")
print("Model:", EMBEDDING_MODEL)
print("Device:", DEVICE)
print("=" * 60)

embedding_model = SentenceTransformer(
    EMBEDDING_MODEL,
    device=DEVICE
)

embedding_model.eval()

print("Embedding model loaded successfully.")
print("=" * 60)


# ============================================================
# SKILL ALIASES
# ============================================================

SKILL_ALIASES = {

    "python": [
        "python",
        "python3"
    ],

    "machine learning": [
        "machine learning",
        "machine-learning",
        "ml"
    ],

    "deep learning": [
        "deep learning",
        "deep-learning",
        "dl"
    ],

    "artificial intelligence": [
        "artificial intelligence",
        "ai"
    ],

    "natural language processing": [
        "natural language processing",
        "nlp"
    ],

    "computer vision": [
        "computer vision",
        "cv"
    ],

    "data analysis": [
        "data analysis",
        "data analytics",
        "data analyst"
    ],

    "data science": [
        "data science",
        "data scientist"
    ],

    "sql": [
        "sql",
        "structured query language"
    ],

    "pandas": [
        "pandas"
    ],

    "numpy": [
        "numpy"
    ],

    "scikit-learn": [
        "scikit-learn",
        "scikit learn",
        "sklearn"
    ],

    "tensorflow": [
        "tensorflow",
        "tf"
    ],

    "keras": [
        "keras"
    ],

    "pytorch": [
        "pytorch",
        "torch"
    ],

    "jupyter": [
        "jupyter",
        "jupyter notebook",
        "notebook"
    ],

    "power bi": [
        "power bi",
        "powerbi"
    ],

    "tableau": [
        "tableau"
    ],

    "excel": [
        "microsoft excel",
        "ms excel",
        "excel"
    ],

    "git": [
        "git",
        "github",
        "gitlab"
    ],

    "docker": [
        "docker"
    ],

    "aws": [
        "aws",
        "amazon web services"
    ],

    "azure": [
        "azure",
        "microsoft azure"
    ],

    "gcp": [
        "gcp",
        "google cloud",
        "google cloud platform"
    ],

    "rest api": [
        "rest api",
        "restful api",
        "restful apis"
    ],

    "flutter": [
        "flutter"
    ],

    "firebase": [
        "firebase"
    ],

    "llm": [
        "llm",
        "large language model",
        "large language models"
    ],

    "rag": [
        "rag",
        "retrieval augmented generation",
        "retrieval-augmented generation"
    ],

    "reinforcement learning": [
        "reinforcement learning",
        "rl"
    ],

    "supervised learning": [
        "supervised learning"
    ],

    "unsupervised learning": [
        "unsupervised learning"
    ],

    "neural networks": [
        "neural network",
        "neural networks",
        "ann",
        "artificial neural network"
    ],

    "data visualization": [
        "data visualization",
        "data visualisation"
    ],

    "matplotlib": [
        "matplotlib"
    ],

    "seaborn": [
        "seaborn"
    ],

    "fastapi": [
        "fastapi"
    ],

    "streamlit": [
        "streamlit"
    ],

    "flask": [
        "flask"
    ],

    "java": [
        "java"
    ],

    "javascript": [
        "javascript",
        "js"
    ],

    "html": [
        "html"
    ],

    "css": [
        "css"
    ]
}


# ============================================================
# TEXT NORMALIZATION
# ============================================================

def normalize_text(text):
    if not text:
        return ""

    text = text.lower()

    text = re.sub(
        r"[\u2010-\u2015]",
        "-",
        text
    )

    text = re.sub(
        r"[ \t]+",
        " ",
        text
    )

    return text.strip()


# ============================================================
# EXTRACT TEXT
# ============================================================

def extract_text(file_path):

    if not file_path:
        raise ValueError("No CV file was provided.")

    extension = os.path.splitext(
        file_path
    )[1].lower()

    # PDF
    if extension == ".pdf":

        pages = []

        with pymupdf.open(file_path) as pdf:

            for page in pdf:

                text = page.get_text("text")

                if text:
                    pages.append(text)

        return "\n".join(pages).strip()

    # DOCX
    if extension == ".docx":

        document = Document(file_path)

        parts = []

        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:
                parts.append(text)

        # Tables
        for table in document.tables:

            for row in table.rows:

                row_text = " ".join(
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                )

                if row_text:
                    parts.append(row_text)

        return "\n".join(parts).strip()

    raise ValueError(
        "Only PDF and DOCX files are supported."
    )


# ============================================================
# BASIC CANDIDATE INFORMATION
# ============================================================

def extract_candidate_name(cv_text):

    lines = [
        line.strip()
        for line in cv_text.splitlines()
        if line.strip()
    ]

    ignored = {
        "resume",
        "cv",
        "curriculum vitae",
        "profile"
    }

    for line in lines[:12]:

        lower = line.lower()

        if lower in ignored:
            continue

        if "@" in line:
            continue

        if re.search(
            r"\b(phone|mobile|email|address)\b",
            lower
        ):
            continue

        if 2 <= len(line.split()) <= 5:

            if len(line) <= 60:

                return line

    return "Unknown Candidate"


def extract_email(cv_text):

    match = re.search(
        r"[A-Za-z0-9._%+-]+"
        r"@[A-Za-z0-9.-]+\."
        r"[A-Za-z]{2,}",
        cv_text
    )

    return match.group(0) if match else ""


def extract_phone(cv_text):

    match = re.search(
        r"\+?\d[\d\s\-\(\)]{8,}\d",
        cv_text
    )

    return match.group(0).strip() if match else ""


# ============================================================
# CV SECTION EXTRACTION
# ============================================================

CV_HEADINGS = [
    "summary",
    "professional summary",
    "profile",
    "objective",
    "skills",
    "technical skills",
    "technical competencies",
    "education",
    "academic background",
    "experience",
    "work experience",
    "professional experience",
    "projects",
    "academic projects",
    "certifications",
    "certificates"
]


def extract_cv_sections(cv_text):

    sections = {}

    lines = [
        line.strip()
        for line in cv_text.splitlines()
        if line.strip()
    ]

    current_section = "general"

    sections[current_section] = []

    for line in lines:

        normalized = normalize_text(line)

        detected_section = None

        for heading in CV_HEADINGS:

            if normalized.rstrip(":") == heading:

                detected_section = heading
                break

        if detected_section:

            current_section = detected_section

            if current_section not in sections:
                sections[current_section] = []

            continue

        sections.setdefault(
            current_section,
            []
        ).append(line)

    return sections


# ============================================================
# CANDIDATE PROFILE
# ============================================================

def extract_candidate_profile(cv_text):

    sections = extract_cv_sections(
        cv_text
    )

    profile = {
        "candidate_name": extract_candidate_name(cv_text),
        "email": extract_email(cv_text),
        "phone": extract_phone(cv_text),
        "education": [],
        "skills": [],
        "experience": [],
        "projects": [],
        "certifications": []
    }

    for key in [
        "education",
        "academic background"
    ]:

        for item in sections.get(key, []):

            profile["education"].append({
                "evidence": item
            })

    for key in [
        "skills",
        "technical skills",
        "technical competencies"
    ]:

        for item in sections.get(key, []):

            profile["skills"].append({
                "skill": item,
                "evidence": item
            })

    for key in [
        "experience",
        "work experience",
        "professional experience"
    ]:

        for item in sections.get(key, []):

            profile["experience"].append({
                "description": item,
                "evidence": item
            })

    for key in [
        "projects",
        "academic projects"
    ]:

        for item in sections.get(key, []):

            profile["projects"].append({
                "description": item,
                "evidence": item
            })

    for key in [
        "certifications",
        "certificates"
    ]:

        for item in sections.get(key, []):

            profile["certifications"].append({
                "name": item,
                "evidence": item
            })

    return profile


# ============================================================
# BUILD CV EVIDENCE
# ============================================================

def build_cv_chunks(cv_text):

    lines = [
        line.strip()
        for line in cv_text.splitlines()
        if line.strip()
    ]

    chunks = []

    current = []
    word_count = 0

    for line in lines:

        words = line.split()

        if (
            word_count + len(words) > 80
            and current
        ):

            chunks.append(
                " ".join(current)
            )

            current = []
            word_count = 0

        current.append(line)
        word_count += len(words)

    if current:

        chunks.append(
            " ".join(current)
        )

    unique = []
    seen = set()

    for chunk in chunks:

        normalized = normalize_text(
            chunk
        )

        if len(normalized) < 10:
            continue

        if normalized not in seen:

            seen.add(normalized)
            unique.append(chunk)

    return unique


# ============================================================
# SKILL MATCHING
# ============================================================

def find_skill_matches(
    requirement,
    candidate_text
):

    requirement_normalized = normalize_text(
        requirement
    )

    candidate_normalized = normalize_text(
        candidate_text
    )

    matched = []

    for canonical, aliases in SKILL_ALIASES.items():

        required_here = any(
            re.search(
                rf"\b{re.escape(alias.lower())}\b",
                requirement_normalized
            )
            for alias in aliases
        )

        if not required_here:
            continue

        candidate_has_skill = any(
            re.search(
                rf"\b{re.escape(alias.lower())}\b",
                candidate_normalized
            )
            for alias in aliases
        )

        if candidate_has_skill:
            matched.append(canonical)

    return matched


def calculate_keyword_score(
    requirement,
    candidate_text
):

    requirement_normalized = normalize_text(
        requirement
    )

    candidate_normalized = normalize_text(
        candidate_text
    )

    matched_skills = find_skill_matches(
        requirement,
        candidate_text
    )

    if matched_skills:
        return 1.0, matched_skills

    if (
        len(requirement_normalized) > 5
        and requirement_normalized
        in candidate_normalized
    ):
        return 1.0, []

    requirement_words = set(
        re.findall(
            r"\b[a-zA-Z][a-zA-Z0-9+#.-]{2,}\b",
            requirement_normalized
        )
    )

    candidate_words = set(
        re.findall(
            r"\b[a-zA-Z][a-zA-Z0-9+#.-]{2,}\b",
            candidate_normalized
        )
    )

    stop_words = {
        "the",
        "and",
        "with",
        "for",
        "from",
        "that",
        "this",
        "are",
        "have",
        "has",
        "using",
        "use",
        "ability",
        "strong",
        "good",
        "excellent",
        "required",
        "experience"
    }

    requirement_words -= stop_words

    if not requirement_words:
        return 0.0, []

    overlap = (
        requirement_words
        & candidate_words
    )

    score = (
        len(overlap)
        / len(requirement_words)
    )

    return min(score, 1.0), []


# ============================================================
# EVIDENCE SCORE
# ============================================================

def calculate_evidence_score(
    keyword_score,
    semantic_score
):

    if keyword_score >= 1.0:
        return 1.0

    if semantic_score >= 0.75:
        return 0.90

    if semantic_score >= 0.65:
        return 0.75

    if semantic_score >= 0.50:
        return 0.60

    if semantic_score >= 0.40:
        return 0.35

    return 0.10


# ============================================================
# JOB TITLE
# ============================================================

def extract_job_title(job_description):

    patterns = [
        r"(?:role|position|job title)\s*:\s*([^\n]+)",
        r"(?:job)\s*:\s*([^\n]+)"
    ]

    for pattern in patterns:

        match = re.search(
            pattern,
            job_description,
            re.IGNORECASE
        )

        if match:
            return match.group(1).strip()

    return "Candidate Position"


# ============================================================
# JOB SECTIONS
# ============================================================

def extract_job_section(
    job_description,
    start_headers,
    end_headers
):

    start_pattern = "|".join(
        re.escape(x)
        for x in start_headers
    )

    end_pattern = "|".join(
        re.escape(x)
        for x in end_headers
    )

    pattern = (
        rf"(?:{start_pattern})"
        rf"\s*:?\s*"
        rf"(.*?)"
        rf"(?=(?:{end_pattern})"
        rf"\s*:|$)"
    )

    match = re.search(
        pattern,
        job_description,
        re.IGNORECASE | re.DOTALL
    )

    return (
        match.group(1).strip()
        if match
        else ""
    )


def split_requirements(text):

    if not text:
        return []

    text = text.replace("•", "\n")
    text = text.replace("●", "\n")
    text = text.replace("▪", "\n")

    requirements = []

    for line in text.splitlines():

        line = line.strip()

        line = re.sub(
            r"^[\-\*\d\.\)\s]+",
            "",
            line
        ).strip()

        if len(line) >= 10:
            requirements.append(line)

    return list(
        dict.fromkeys(requirements)
    )


# ============================================================
# CATEGORY DETECTION
# ============================================================

def detect_category(requirement):

    text = normalize_text(
        requirement
    )

    if any(
        x in text
        for x in [
            "bachelor",
            "master",
            "degree",
            "education",
            "university",
            "college"
        ]
    ):
        return "education"

    if any(
        x in text
        for x in [
            "year experience",
            "years experience",
            "experience",
            "worked",
            "background"
        ]
    ):
        return "experience"

    if any(
        x in text
        for x in [
            "communication",
            "leadership",
            "teamwork",
            "collaboration",
            "problem solving",
            "analytical"
        ]
    ):
        return "soft_skill"

    for aliases in SKILL_ALIASES.values():

        if any(
            alias in text
            for alias in aliases
        ):
            return "skill"

    return "knowledge"


# ============================================================
# JOB REQUIREMENTS
# ============================================================

def extract_job_requirements(
    job_description
):

    text = job_description.strip()

    job_title = extract_job_title(text)

    required_section = extract_job_section(
        text,
        [
            "Required Skills & Qualifications",
            "Required Qualifications",
            "Required Skills",
            "Requirements"
        ],
        [
            "Preferred Skills",
            "Preferred Qualifications",
            "Preferred",
            "More About",
            "Equal Opportunity",
            "Benefits"
        ]
    )

    required_items = split_requirements(
        required_section
    )

    preferred_section = extract_job_section(
        text,
        [
            "Preferred Skills",
            "Preferred Qualifications",
            "Preferred"
        ],
        [
            "More About",
            "Equal Opportunity",
            "Benefits",
            "Responsibilities",
            "Key Responsibilities"
        ]
    )

    preferred_items = split_requirements(
        preferred_section
    )

    if not required_items:

        required_items = split_requirements(
            text
        )

    required = []

    if required_items:

        weight = 85.0 / len(
            required_items
        )

        for item in required_items:

            required.append({
                "requirement": item,
                "category": detect_category(item),
                "weight": round(weight, 2),
                "type": "required"
            })

    preferred = []

    if preferred_items:

        weight = 15.0 / len(
            preferred_items
        )

        for item in preferred_items:

            preferred.append({
                "requirement": item,
                "category": detect_category(item),
                "weight": round(weight, 2),
                "type": "preferred"
            })

    responsibilities_section = extract_job_section(
        text,
        [
            "Key Responsibilities",
            "Responsibilities",
            "Role Overview"
        ],
        [
            "Required Skills",
            "Required Qualifications",
            "Preferred",
            "More About",
            "Equal Opportunity"
        ]
    )

    responsibilities = split_requirements(
        responsibilities_section
    )

    return {
        "job_title": job_title,
        "required": required,
        "preferred": preferred,
        "responsibilities": responsibilities
    }


# ============================================================
# MATCH REQUIREMENTS
# ============================================================

def match_requirements(
    requirements,
    candidate_chunks
):

    if not requirements:
        return []

    if not candidate_chunks:

        return [
            {
                "requirement": req["requirement"],
                "category": req["category"],
                "weight": req["weight"],
                "similarity": 0.0,
                "keyword_score": 0.0,
                "evidence_score": 0.0,
                "final_score": 0.0,
                "status": "Weak / No Match",
                "evidence": "No CV evidence found.",
                "matched_skills": []
            }
            for req in requirements
        ]

    # Encode CV ONCE
    candidate_embeddings = embedding_model.encode(
        candidate_chunks,
        normalize_embeddings=True,
        convert_to_numpy=True,
        show_progress_bar=False,
        batch_size=32
    )

    # Encode requirements ONCE
    requirement_texts = [
        req["requirement"]
        for req in requirements
    ]

    requirement_embeddings = embedding_model.encode(
        requirement_texts,
        normalize_embeddings=True,
        convert_to_numpy=True,
        show_progress_bar=False,
        batch_size=32
    )

    # Matrix multiplication because embeddings
    # are already normalized.
    similarity_matrix = np.matmul(
        requirement_embeddings,
        candidate_embeddings.T
    )

    results = []

    for index, req in enumerate(requirements):

        similarities = similarity_matrix[
            index
        ]

        best_index = int(
            np.argmax(similarities)
        )

        semantic_score = float(
            similarities[best_index]
        )

        evidence = candidate_chunks[
            best_index
        ]

        keyword_score, matched_skills = (
            calculate_keyword_score(
                req["requirement"],
                evidence
            )
        )

        # Search all chunks for explicit keyword evidence.
        for chunk in candidate_chunks:

            current_score, current_skills = (
                calculate_keyword_score(
                    req["requirement"],
                    chunk
                )
            )

            if current_score > keyword_score:

                keyword_score = current_score
                matched_skills = current_skills
                evidence = chunk

        evidence_score = (
            calculate_evidence_score(
                keyword_score,
                semantic_score
            )
        )

        final_score = (
            SEMANTIC_WEIGHT * semantic_score
            +
            KEYWORD_WEIGHT * keyword_score
            +
            EVIDENCE_WEIGHT * evidence_score
        )

        final_score = max(
            0.0,
            min(1.0, final_score)
        )

        if final_score >= STRONG_MATCH:

            status = "Strong Match"

        elif final_score >= PARTIAL_MATCH:

            status = "Partial Match"

        else:

            status = "Weak / No Match"

        results.append({
            "requirement": req["requirement"],
            "category": req["category"],
            "weight": req["weight"],
            "type": req.get(
                "type",
                "required"
            ),
            "similarity": round(
                semantic_score,
                3
            ),
            "keyword_score": round(
                keyword_score,
                3
            ),
            "evidence_score": round(
                evidence_score,
                3
            ),
            "final_score": round(
                final_score,
                3
            ),
            "status": status,
            "evidence": evidence,
            "matched_skills": matched_skills
        })

    return results


# ============================================================
# SCORE
# ============================================================

def calculate_score(results):

    if not results:
        return 0.0

    total_weight = sum(
        float(x["weight"])
        for x in results
    )

    if total_weight <= 0:
        return 0.0

    weighted_score = sum(
        x["final_score"] * x["weight"]
        for x in results
    )

    return round(
        weighted_score
        / total_weight
        * 100,
        2
    )


# ============================================================
# RECOMMENDATION
# ============================================================

def generate_recommendation(
    score,
    required_coverage
):

    if (
        score >= 80
        and required_coverage >= 75
    ):

        return (
            "SHORTLIST",
            "Strong overall match with good coverage of required qualifications."
        )

    if (
        score >= 70
        and required_coverage >= 65
    ):

        return (
            "SHORTLIST",
            "Good overall match with substantial evidence for required qualifications."
        )

    if (
        score >= 55
        and required_coverage >= 60
    ):

        return (
            "REVIEW",
            "The candidate shows meaningful alignment with the requirements; human review is recommended."
        )

    if required_coverage >= 75:

        return (
            "REVIEW",
            "Most required qualifications appear to be covered despite a moderate semantic score."
        )

    return (
        "DO NOT SHORTLIST",
        "Insufficient evidence of alignment with the required qualifications."
    )


# ============================================================
# MAIN SCREENING FUNCTION
# ============================================================

def screen_candidate(
    cv_path,
    job_description
):

    total_start = time.perf_counter()

    # --------------------------------------------------------
    # CV extraction
    # --------------------------------------------------------

    start = time.perf_counter()

    cv_text = extract_text(
        cv_path
    )

    if not cv_text.strip():

        raise ValueError(
            "No readable text was found in the CV."
        )

    extraction_time = (
        time.perf_counter()
        - start
    )

    # --------------------------------------------------------
    # Candidate profile
    # --------------------------------------------------------

    start = time.perf_counter()

    candidate_profile = (
        extract_candidate_profile(
            cv_text
        )
    )

    profile_time = (
        time.perf_counter()
        - start
    )

    # --------------------------------------------------------
    # Job requirements
    # --------------------------------------------------------

    start = time.perf_counter()

    job_result = (
        extract_job_requirements(
            job_description
        )
    )

    job_time = (
        time.perf_counter()
        - start
    )

    # --------------------------------------------------------
    # CV evidence
    # --------------------------------------------------------

    start = time.perf_counter()

    candidate_chunks = build_cv_chunks(
        cv_text
    )

    evidence_time = (
        time.perf_counter()
        - start
    )

    # --------------------------------------------------------
    # Semantic matching
    # --------------------------------------------------------

    start = time.perf_counter()

    requirements = (
        job_result["required"]
        +
        job_result["preferred"]
    )

    matching_results = (
        match_requirements(
            requirements,
            candidate_chunks
        )
    )

    matching_time = (
        time.perf_counter()
        - start
    )

    # --------------------------------------------------------
    # Score
    # --------------------------------------------------------

    start = time.perf_counter()

    overall_score = calculate_score(
        matching_results
    )

    required_results = [
        x
        for x in matching_results
        if x["type"] == "required"
    ]

    if required_results:

        covered = sum(
            1
            for x in required_results
            if x["final_score"] >= 0.50
        )

        required_coverage = round(
            covered
            / len(required_results)
            * 100,
            2
        )

    else:

        required_coverage = 0.0

    recommendation, reason = (
        generate_recommendation(
            overall_score,
            required_coverage
        )
    )

    scoring_time = (
        time.perf_counter()
        - start
    )

    total_time = (
        time.perf_counter()
        - total_start
    )

    # --------------------------------------------------------
    # Performance log
    # --------------------------------------------------------

    print("\n" + "=" * 60)
    print("SCREENING PERFORMANCE")
    print("=" * 60)

    print(
        f"CV extraction:       {extraction_time:.2f}s"
    )

    print(
        f"Candidate profile:   {profile_time:.2f}s"
    )

    print(
        f"Job analysis:        {job_time:.2f}s"
    )

    print(
        f"CV evidence:         {evidence_time:.2f}s"
    )

    print(
        f"Semantic matching:   {matching_time:.2f}s"
    )

    print(
        f"Scoring:             {scoring_time:.2f}s"
    )

    print("-" * 60)

    print(
        f"TOTAL SCREENING:     {total_time:.2f}s"
    )

    print(
        f"OVERALL SCORE:       {overall_score:.2f}"
    )

    print(
        f"REQUIRED COVERAGE:   {required_coverage:.2f}%"
    )

    print("=" * 60)

    return {
        "candidate": candidate_profile,
        "job": job_result,
        "score": overall_score,
        "required_coverage": required_coverage,
        "recommendation": recommendation,
        "reason": reason,
        "human_review_required": True,
        "matches": matching_results,
        "processing_time": round(
            total_time,
            2
        )
    }
